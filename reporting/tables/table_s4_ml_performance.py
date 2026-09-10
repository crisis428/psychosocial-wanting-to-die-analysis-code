from __future__ import annotations

import argparse
import json
import zipfile
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

CSV_MEMBER = "combined/supplement_table_detailed_performance.csv"
GROUP_ORDER = ["total", "male", "female"]
MODEL_ORDER = [
    "LogisticRegression", "RandomForest", "DecisionTree",
    "SVM", "XGBoost", "CatBoost",
]
GROUP_LABEL = {
    "total": "Total Sample",
    "male": "Male Subgroup",
    "female": "Female Subgroup",
}
MODEL_LABEL = {
    "LogisticRegression": "Logistic Regression",
    "RandomForest": "Random Forest",
    "DecisionTree": "Decision Tree",
    "SVM": "SVM",
    "XGBoost": "XGBoost",
    "CatBoost": "CatBoost",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate final expanded Table S4.")
    parser.add_argument(
        "--results-zip", type=Path, default=Path("full_results_checkpointed.zip")
    )
    parser.add_argument("--output-dir", type=Path, default=Path("table_s4_ml_performance"))
    return parser.parse_args()


def format_ci(estimate: float, lower: float, upper: float) -> str:
    return f"{estimate:.3f} ({lower:.3f}–{upper:.3f})"


def load_performance(results_zip: Path) -> pd.DataFrame:
    with zipfile.ZipFile(results_zip) as archive:
        member = CSV_MEMBER
        if member not in archive.namelist():
            candidates = [
                name for name in archive.namelist()
                if name.endswith("supplement_table_detailed_performance.csv")
            ]
            if len(candidates) != 1:
                raise FileNotFoundError(candidates)
            member = candidates[0]
        with archive.open(member) as handle:
            frame = pd.read_csv(handle)

    required = {
        "Group", "Model", "N", "Events", "AUROC", "AUROC_CI_Low",
        "AUROC_CI_High", "Average_Precision", "AP_CI_Low", "AP_CI_High",
        "Brier", "Brier_CI_Low", "Brier_CI_High", "Log_Loss",
        "Calibration_Intercept", "Calibration_Slope", "Sensitivity",
        "Specificity", "Precision", "Accuracy", "F1", "Threshold",
    }
    missing = required - set(frame.columns)
    if missing:
        raise ValueError(f"Missing columns: {sorted(missing)}")
    if len(frame) != 18:
        raise ValueError(f"Expected 18 rows; found {len(frame)}")
    if not (frame["Threshold"].round(10) == 0.50).all():
        raise ValueError("Threshold is not 0.50 in all rows")
    return frame


def build_panels(frame: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    frame = frame.copy()
    frame["_group_order"] = frame["Group"].map({v: i for i, v in enumerate(GROUP_ORDER)})
    frame["_model_order"] = frame["Model"].map({v: i for i, v in enumerate(MODEL_ORDER)})
    frame = frame.sort_values(["_group_order", "_model_order"]).reset_index(drop=True)
    frame["Group_Label"] = frame["Group"].map(GROUP_LABEL)
    frame["Model_Label"] = frame["Model"].map(MODEL_LABEL)

    panel_a = pd.DataFrame({
        "Group": frame["Group"],
        "Group Label": frame["Group_Label"],
        "N": frame["N"].astype(int),
        "Events": frame["Events"].astype(int),
        "Model": frame["Model_Label"],
        "AUROC (95% CI)": [format_ci(a, b, c) for a, b, c in zip(frame.AUROC, frame.AUROC_CI_Low, frame.AUROC_CI_High)],
        "Average Precision (95% CI)": [format_ci(a, b, c) for a, b, c in zip(frame.Average_Precision, frame.AP_CI_Low, frame.AP_CI_High)],
        "Brier Score (95% CI)": [format_ci(a, b, c) for a, b, c in zip(frame.Brier, frame.Brier_CI_Low, frame.Brier_CI_High)],
        "Log Loss": frame.Log_Loss.map(lambda value: f"{value:.3f}"),
        "Calibration Intercept": frame.Calibration_Intercept.map(lambda value: f"{value:.3f}"),
        "Calibration Slope": frame.Calibration_Slope.map(lambda value: f"{value:.3f}"),
    })
    panel_b = pd.DataFrame({
        "Group": frame["Group"],
        "Group Label": frame["Group_Label"],
        "N": frame["N"].astype(int),
        "Events": frame["Events"].astype(int),
        "Model": frame["Model_Label"],
        "Sensitivity": frame.Sensitivity.map(lambda value: f"{value:.3f}"),
        "Specificity": frame.Specificity.map(lambda value: f"{value:.3f}"),
        "Precision": frame.Precision.map(lambda value: f"{value:.3f}"),
        "Accuracy": frame.Accuracy.map(lambda value: f"{value:.3f}"),
        "F1 Score": frame.F1.map(lambda value: f"{value:.3f}"),
    })
    return panel_a, panel_b


def write_excel(panel_a: pd.DataFrame, panel_b: pd.DataFrame, path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Table S4"
    sheet.sheet_view.showGridLines = False

    dark_fill = PatternFill("solid", fgColor="1F4E78")
    header_fill = PatternFill("solid", fgColor="D9E2F3")
    group_fill = PatternFill("solid", fgColor="EDEDED")
    thin = Side(style="thin", color="B7B7B7")
    medium = Side(style="medium", color="000000")

    row = 1
    title = "Table S4. Performance Comparison of 6 Models for Classifying Past-Year Thoughts of Wanting to Die"
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
    cell = sheet.cell(row=row, column=1, value=title)
    cell.font = Font(bold=True, size=12)
    cell.alignment = Alignment(wrap_text=True)
    row += 2

    def add_panel(title_text: str, frame: pd.DataFrame, columns: list[str]) -> int:
        nonlocal row
        width = len(columns)
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
        panel_cell = sheet.cell(row=row, column=1, value=title_text)
        panel_cell.fill = dark_fill
        panel_cell.font = Font(bold=True, color="FFFFFF", size=10)
        row += 1

        for col_index, header in enumerate(columns, start=1):
            header_cell = sheet.cell(row=row, column=col_index, value=header)
            header_cell.fill = header_fill
            header_cell.font = Font(bold=True, size=9)
            header_cell.alignment = Alignment(horizontal="center", wrap_text=True)
            header_cell.border = Border(top=medium, bottom=thin)
        row += 1

        for group in GROUP_ORDER:
            group_frame = frame.loc[frame["Group"] == group]
            n_value = int(group_frame["N"].iloc[0])
            events = int(group_frame["Events"].iloc[0])
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
            group_cell = sheet.cell(
                row=row, column=1,
                value=f"{GROUP_LABEL[group]} (N = {n_value:,}; events = {events:,})",
            )
            group_cell.fill = group_fill
            group_cell.font = Font(bold=True, size=9)
            row += 1

            for _, result in group_frame.iterrows():
                for col_index, column in enumerate(columns, start=1):
                    value = result[column]
                    body_cell = sheet.cell(row=row, column=col_index, value=value)
                    body_cell.font = Font(size=9)
                    body_cell.alignment = Alignment(
                        horizontal="left" if col_index == 1 else "center",
                        vertical="center", wrap_text=True,
                    )
                    body_cell.border = Border(bottom=thin)
                row += 1
            for col_index in range(1, width + 1):
                body_cell = sheet.cell(row=row - 1, column=col_index)
                body_cell.border = Border(bottom=medium)
            row += 1
        return width

    width_a = add_panel(
        "Panel A. Discrimination, Overall Performance, and Calibration",
        panel_a,
        ["Model", "AUROC (95% CI)", "Average Precision (95% CI)", "Brier Score (95% CI)", "Log Loss", "Calibration Intercept", "Calibration Slope"],
    )
    width_b = add_panel(
        "Panel B. Threshold-Based Performance at a Probability Threshold of 0.50",
        panel_b,
        ["Model", "Sensitivity", "Specificity", "Precision", "Accuracy", "F1 Score"],
    )

    footnote = (
        "AUROC, average precision, and Brier score are presented with 95% CIs estimated from 2000 participant-level "
        "bootstrap replicates. Threshold-based metrics were calculated at a probability threshold of 0.50. Values are "
        "based on participant-level aggregated out-of-fold predictions from repeated nested cross-validation."
    )
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max(width_a, width_b))
    footnote_cell = sheet.cell(row=row, column=1, value=footnote)
    footnote_cell.font = Font(size=8, italic=True)
    footnote_cell.alignment = Alignment(wrap_text=True)

    for index, width in {1: 24, 2: 22, 3: 27, 4: 23, 5: 14, 6: 21, 7: 18}.items():
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins = PageMargins(left=0.25, right=0.25, top=0.35, bottom=0.35)
    workbook.save(path)


def main() -> None:
    args = parse_args()
    results_zip = args.results_zip.resolve()
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    frame = load_performance(results_zip)
    panel_a, panel_b = build_panels(frame)
    panel_a.to_csv(output_dir / "Table_S4_Panel_A.csv", index=False, encoding="utf-8-sig")
    panel_b.to_csv(output_dir / "Table_S4_Panel_B.csv", index=False, encoding="utf-8-sig")
    write_excel(panel_a, panel_b, output_dir / "Table_S4_ML_Performance.xlsx")
    (output_dir / "generation_manifest.json").write_text(
        json.dumps({"source": str(results_zip), "rows": len(frame)}, indent=2),
        encoding="utf-8",
    )
    print(f"Completed: {output_dir}")


if __name__ == "__main__":
    main()
